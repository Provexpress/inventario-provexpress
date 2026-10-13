"""
═══════════════════════════════════════════════════════════════════════════════════
PROVEXPRESS - ACTUALIZADOR DIARIO DE LA HOJA 'GENERADOR CORREO INVENTARIO'
═══════════════════════════════════════════════════════════════════════════════════
Este script sincroniza automáticamente el Excel de OneDrive:
'plantilla_inventario_matriz.xlsx'
Garantizando que la hoja 'Generador Correo Inventario' sea un reflejo EXACTO 1:1
del correo de inventario diario enviado por Alejandro Varón en Outlook.
═══════════════════════════════════════════════════════════════════════════════════
"""

import os
import sys
import re
import json
import time
import traceback
from datetime import datetime
from bs4 import BeautifulSoup
import openpyxl

# --- RUTAS DE ARCHIVOS ---
EXCEL_PATH = r"C:\Users\especialista.prevent\OneDrive - PROVEXPRESS SAS\Excels Invetario\plantilla_inventario_matriz.xlsx"
LOG_DIR = r"C:\Users\especialista.prevent\OneDrive - PROVEXPRESS SAS\Excels Invetario"
LOG_FILE_ONEDRIVE = os.path.join(LOG_DIR, "registro_actualizaciones_inventario.log")
LOG_FILE_LOCAL = r"c:\Proyectos\inventario-provexpress\actualizar_inventario.log"

PROJECT_PATHS = [
    r"c:\Proyectos\inventario-provexpress",
    r"C:\Users\especialista.prevent\OneDrive - PROVEXPRESS SAS\Documents\Proyectos\inventario-provexpress"
]

MESES_ESPANOL = {
    1: "ENERO", 2: "FEBRERO", 3: "MARZO", 4: "ABRIL",
    5: "MAYO", 6: "JUNIO", 7: "JULIO", 8: "AGOSTO",
    9: "SEPTIEMBRE", 10: "OCTUBRE", 11: "NOVIEMBRE", 12: "DICIEMBRE"
}

def log_message(msg):
    ts = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    full_msg = f"[{ts}] {msg}"
    try:
        if sys.stdout and not sys.stdout.closed:
            print(full_msg, flush=True)
    except Exception:
        pass
        
    for log_p in [LOG_FILE_ONEDRIVE, LOG_FILE_LOCAL]:
        try:
            os.makedirs(os.path.dirname(log_p), exist_ok=True)
            with open(log_p, "a", encoding="utf-8") as f:
                f.write(full_msg + "\n")
                f.flush()
        except Exception:
            pass

def parse_colombian_price(cost_text):
    part0 = cost_text.split('/')[0]
    clean = re.sub(r'[^\d,\.]', '', part0).strip('.,')
    if not clean: return 0.0
    if ',' in clean:
        parts_comma = clean.split(',')
        if len(parts_comma) == 2 and len(parts_comma[1]) <= 2:
            int_part = parts_comma[0].replace('.', '')
            return float(f"{int_part}.{parts_comma[1]}")
        else:
            clean = clean.replace(',', '')
    if '.' in clean:
        parts_dot = clean.split('.')
        if len(parts_dot[-1]) <= 2 and len(parts_dot) > 1:
            int_part = ''.join(parts_dot[:-1])
            dec_part = parts_dot[-1]
            return float(f"{int_part}.{dec_part}")
        else:
            return float(''.join(parts_dot))
    try:
        return float(clean)
    except Exception:
        return 0.0

def parse_email_tables(html_content, cid_to_file=None):
    if cid_to_file is None:
        cid_to_file = {}
        
    soup = BeautifulSoup(html_content, 'html.parser')
    wrapper = soup.body.find('div')
    if not wrapper:
        wrapper = soup.body
    if not wrapper:
        return []
    
    tables = wrapper.find_all('table')
    products = []
    
    def get_preceding_section(table):
        prev = table.find_previous(
            lambda tag: tag.name in ['p', 'div', 'h1', 'h2', 'h3', 'b', 'strong'] and
            any(s in tag.get_text().upper() for s in ['RENTA', 'HP', 'LENOVO', 'DELL', 'APPLE', 'MONITOR', 'IMPRESORA', 'CELULAR', 'VENTA']) and
            len(tag.get_text().strip()) < 40 and not tag.find_parent('table')
        )
        return prev.get_text().strip() if prev else 'VENTA'

    for t in tables:
        raw_text = t.get_text(separator=' | ', strip=True)
        rows = t.find_all('tr')
        
        np_val = ''
        prod_desc = ''
        cost_text = ''
        ingreso_text = ''
        img_filename = ''
        
        img_tag = t.find('img')
        if img_tag and img_tag.get('src'):
            src = img_tag['src']
            m_cid = re.search(r'cid:([^@"\'>\s]+)', src, re.I)
            if m_cid:
                img_filename = m_cid.group(1).lower()

        for r in rows:
            tds = r.find_all(['td', 'th'])
            for i, td in enumerate(tds):
                txt = td.get_text(strip=True).upper()
                if ('N/P' in txt or 'N/PARTE' in txt or 'PARTE' in txt) and i + 1 < len(tds):
                    np_val = tds[i+1].get_text(strip=True)
                elif ('PRODUCTO' in txt or 'NOMBRE' in txt or 'DESCRIPCION' in txt) and i + 1 < len(tds):
                    prod_desc = tds[i+1].get_text(strip=True)
                elif 'COSTO' in txt and i + 1 < len(tds):
                    cost_text = tds[i+1].get_text(strip=True)
                elif 'INGRESO' in txt:
                    ingreso_text = txt.replace('INGRESO', '').strip()
                elif i + 1 < len(tds) and 'INGRESO' in tds[i+1].get_text().upper():
                    ingreso_text = tds[i+1].get_text().replace('INGRESO', '').replace('ingreso', '').strip()

        if not ingreso_text:
            m_ing = re.search(r'INGRESO\s*(\d{2}/\d{2}/\d{2,4})', raw_text, re.I)
            if m_ing:
                ingreso_text = m_ing.group(1)

        if not (np_val and prod_desc):
            continue

        iva = "Sí" if "+ IVA" in cost_text.upper() or "+IVA" in cost_text.upper() or "+ iva" in cost_text.lower() else "No"
        qty_match = re.search(r'/\s*(\d+)\s*UNIDAD', cost_text, re.I)
        qty = int(qty_match.group(1)) if qty_match else 0
        cost_val = parse_colombian_price(cost_text)

        sec_heading = get_preceding_section(t).upper()

        # Marca precisa
        brand = "OTRAS"
        pd_upper = prod_desc.upper()
        np_upper = np_val.upper()
        
        for b in ["OSTER", "TARGUS", "GENIUS", "LOGITECH", "KYOCERA", "MOTOROLA", "ADATA", "JALTECH", "XKIM", "CISCO", "ARUBA", "RICOH", "POLY", "SAMSUNG", "APPLE", "DELL", "HP", "LENOVO"]:
            if b in pd_upper or b in np_upper:
                brand = b
                break
        if brand == "OTRAS":
            for b in ["HP", "LENOVO", "DELL", "APPLE", "SAMSUNG"]:
                if b in sec_heading:
                    brand = b
                    break
                
        # Categoría Hardware
        cat = "ACCESORIOS Y PERIFÉRICOS"
        if "RENTA" in sec_heading or "RENTA" in pd_upper:
            cat = "EQUIPOS PARA RENTA"
        elif "MORRAL" in pd_upper:
            cat = "ACCESORIOS / MORRALES"
        elif "PORTATIL" in pd_upper or "PROBOOK" in pd_upper or "THINKPAD" in pd_upper or "MACBOOK" in pd_upper:
            cat = "PORTÁTILES"
        elif "MONITOR" in pd_upper or "S3 ESSENTIAL" in pd_upper:
            cat = "MONITORES"
        elif "IMPRESORA" in pd_upper or "MULTIFUNCIONAL" in pd_upper or "MFP" in pd_upper:
            cat = "IMPRESORAS"
        elif "DIADEMA" in pd_upper or "AURICULAR" in pd_upper or "HEADSET" in pd_upper:
            cat = "DIADEMAS / AUDIO"
        elif "MOUSE" in pd_upper or "TECLADO" in pd_upper or "COMBO" in pd_upper:
            cat = "PERIFÉRICOS / MOUSE & TECLADO"
        elif "CELULAR" in pd_upper or "IPAD" in pd_upper or "TABLET" in pd_upper or "MOTO G" in pd_upper:
            cat = "CELULARES / MÓVILES"
        elif "DISCO" in pd_upper or "SSD" in pd_upper or "ALMACENAMIENTO" in pd_upper:
            cat = "ALMACENAMIENTO / SSD"
        elif "BASE" in pd_upper or "REFRIGERANTE" in pd_upper:
            cat = "ACCESORIOS / BASES REFRIGERANTES"
        elif "MORRAL" in pd_upper:
            cat = "ACCESORIOS / MORRALES"
        elif "CABLE" in pd_upper:
            cat = "CABLES / CONECTIVIDAD"
        elif "TARJETA" in pd_upper:
            cat = "TARJETAS DE CONTROL"
        elif "WS" in pd_upper or "THINKSTATION" in pd_upper or "WORKSTATION" in pd_upper:
            cat = "WORKSTATIONS"
        elif "PC" in pd_upper or "DESKTOP" in pd_upper or "CPU HP" in pd_upper:
            cat = "DESKTOPS / PCS"

        # Sección Correo
        sec = "ACCESORIOS Y PERIFÉRICOS"
        if cat == "EQUIPOS PARA RENTA" or "RENTA" in sec_heading:
            sec = "EQUIPOS PARA RENTA"
        elif brand == "HP" and cat in ["PORTÁTILES", "DESKTOPS / PCS", "WORKSTATIONS"]:
            sec = "HP"
        elif brand == "LENOVO" and cat in ["PORTÁTILES", "DESKTOPS / PCS", "WORKSTATIONS"]:
            sec = "LENOVO"
        elif brand == "DELL" and cat in ["PORTÁTILES", "DESKTOPS / PCS", "WORKSTATIONS"]:
            sec = "DELL"
        elif cat == "MONITORES":
            sec = "MONITORES"
        elif cat == "IMPRESORAS":
            sec = "IMPRESORAS"
        elif cat == "CELULARES / MÓVILES":
            sec = "CELULARES"

        local_img = cid_to_file.get(img_filename, "")
        if not local_img:
            if cat == "PORTÁTILES": local_img = "https://images.unsplash.com/photo-1588872657578-7efd1f1555ed?w=300"
            elif cat == "MONITORES": local_img = "https://images.unsplash.com/photo-1527443224154-c4a3942d3acf?w=300"
            elif cat == "IMPRESORAS": local_img = "https://images.unsplash.com/photo-1612815154858-60aa4c59eaa6?w=300"
            elif cat == "DIADEMAS / AUDIO": local_img = "https://images.unsplash.com/photo-1546435770-a3e426bf472b?w=300"
            else: local_img = "https://images.unsplash.com/photo-1519389950473-47ba0277781c?w=300"

        # Se guardan en el orden natural en que aparecen en el correo
        products.append({
            "seccion": sec,
            "categoria": cat,
            "marca": brand,
            "np": np_val,
            "producto": prod_desc,
            "costo": cost_val,
            "iva": iva,
            "fecha_ingreso": ingreso_text,
            "cantidad_actual": qty,
            "imagen": local_img
        })

    return products

def buscar_correos_inventario(outlook_ns):
    root = outlook_ns.Folders.Item('especialista.preventa@provexpress.com.co')
    candidate_folders = []
    
    try:
        candidate_folders.append(root.Folders.Item('INVENTARIOS'))
    except Exception:
        pass
        
    try:
        candidate_folders.append(root.Folders.Item('Bandeja de entrada'))
    except Exception:
        pass

    pattern = re.compile(r'INVENTARIO\s+(\d{1,2})\s+([A-ZÁÉÍÓÚ]+)(?:\s+(\d{4}))?', re.IGNORECASE)

    encontrados = {}
    
    for f in candidate_folders:
        try:
            items = f.Items
            items.Sort("[ReceivedTime]", True)
            limit = min(30, items.Count)
            for i in range(1, limit + 1):
                try:
                    msg = items.Item(i)
                    subj = str(getattr(msg, 'Subject', '') or '')
                    if not subj or 'INVENTARIO' not in subj.upper():
                        continue
                    
                    m = pattern.search(subj)
                    if m:
                        day = int(m.group(1))
                        month_str = m.group(2).upper()
                        rcv = msg.ReceivedTime
                        year = int(m.group(3)) if m.group(3) else rcv.year
                        key = (year, month_str, day)
                        
                        if key not in encontrados or rcv > encontrados[key]['received_time']:
                            encontrados[key] = {
                                'day': day,
                                'month_str': month_str,
                                'year': year,
                                'subject': subj.strip(),
                                'received_time': rcv,
                                'msg': msg
                            }
                except Exception:
                    continue
        except Exception as e:
            log_message(f"Aviso al escanear carpeta '{getattr(f, 'Name', 'desconocida')}': {e}")

    return list(encontrados.values())

def actualizar_excel_y_sistema(force=False):
    log_message("=== INICIANDO SINCRONIZACIÓN DIARIA DE INVENTARIO ===")
    
    if not os.path.exists(EXCEL_PATH):
        log_message(f"ERROR CRÍTICO: No se encontró el archivo Excel en {EXCEL_PATH}")
        return False

    import win32com.client as win32
    try:
        outlook = win32.Dispatch('Outlook.Application').GetNamespace('MAPI')
    except Exception as e:
        log_message(f"ERROR: No se pudo conectar con Outlook: {e}")
        return False

    correos = buscar_correos_inventario(outlook)
    if not correos:
        log_message("No se encontraron correos de inventario en Outlook.")
        return False

    # Ordenar por fecha recibida descendente
    correos.sort(key=lambda x: x['received_time'], reverse=True)
    latest_correo = correos[0]
    
    now = datetime.now()
    curr_month_str = MESES_ESPANOL.get(now.month, "OCTUBRE")
    
    log_message(f"Último correo detectado: '{latest_correo['subject']}' del {latest_correo['day']} de {latest_correo['month_str']} {latest_correo['year']} ({latest_correo['received_time']})")

    # Verificar si el Excel ya tiene este correo aplicado exactamente
    try:
        wb_check = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
        ws_check = wb_check["Generador Correo Inventario"]
        current_subject = str(ws_check["F3"].value or ws_check["E3"].value or "").strip()
        current_status = str(ws_check["J3"].value or ws_check["G3"].value or "").strip()
        wb_check.close()
        
        if not force and current_subject.upper() == latest_correo['subject'].upper() and "ACTUALIZADO" in current_status.upper():
            log_message(f"✓ El Excel ya se encuentra al día con '{latest_correo['subject']}'. No requiere cambios.")
            return True
    except Exception as e:
        log_message(f"Aviso al verificar estado previo del Excel: {e}")

    # Extraer imágenes del correo más reciente
    cid_to_file = {}
    img_dirs = [os.path.join(p, "img", "products") for p in PROJECT_PATHS if os.path.exists(p)]
    for d in img_dirs:
        os.makedirs(d, exist_ok=True)

    today_msg = latest_correo['msg']
    try:
        for i in range(1, today_msg.Attachments.Count + 1):
            att = today_msg.Attachments.Item(i)
            fn = str(att.FileName)
            if any(fn.lower().endswith(ext) for ext in ['.jpg', '.jpeg', '.png', '.webp', '.bmp']):
                for d in img_dirs:
                    save_path = os.path.join(d, fn)
                    try:
                        att.SaveAsFile(save_path)
                    except Exception:
                        pass
                cid_to_file[fn.lower()] = f"img/products/{fn}"
    except Exception as e:
        log_message(f"Aviso al extraer imágenes adjuntas: {e}")

    # Parsear productos del correo en su ORDEN EXACTO ORIGINAL (sin reordenamiento artificial)
    today_products = parse_email_tables(today_msg.HTMLBody, cid_to_file)
    if not today_products:
        log_message("⚠ No se pudieron extraer tablas de productos del correo.")
        return False

    log_message(f"✓ Se extrajeron {len(today_products)} productos vigentes en el orden 1:1 del correo.")

    # Obtener el histórico de días de stock para este mes
    same_month_emails = [c for c in correos if c['month_str'] == latest_correo['month_str'] and c['year'] == latest_correo['year']]
    daily_stock_by_np = {}
    for c in same_month_emails:
        day_num = c['day']
        try:
            p_day = parse_email_tables(c['msg'].HTMLBody, cid_to_file)
            for p in p_day:
                np_k = p['np']
                if np_k not in daily_stock_by_np:
                    daily_stock_by_np[np_k] = {}
                daily_stock_by_np[np_k][f"day_{day_num}"] = p['cantidad_actual']
        except Exception:
            pass

    for p in today_products:
        p['dias_stock'] = daily_stock_by_np.get(p['np'], {f"day_{latest_correo['day']}": p['cantidad_actual']})

    # --- ESCRIBIR EN EXCEL EN ONEDRIVE ---
    log_message(f"Escribiendo en Excel: {EXCEL_PATH} ...")
    
    wb = None
    for intento in range(5):
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH, data_only=False)
            break
        except PermissionError:
            log_message(f"Archivo bloqueado (abierto por usuario). Reintentando en 3s... (Intento {intento+1}/5)")
            time.sleep(3)
        except Exception as e:
            log_message(f"Error abriendo Excel: {e}")
            return False

    if wb is None:
        log_message("ERROR: No se pudo obtener acceso al archivo Excel. Está abierto en Microsoft Excel.")
        return False

    ws_matriz = wb["Matriz Inventario"]
    ws_gen = wb["Generador Correo Inventario"]

    # 1. ACTUALIZAR HOJA 'GENERADOR CORREO INVENTARIO'
    email_date = datetime(latest_correo['year'], now.month if latest_correo['month_str'] == curr_month_str else 10, latest_correo['day'])
    
    # Encabezado (Fila 3)
    ws_gen["B3"].value = "Fecha del Reporte:"
    ws_gen["C3"].value = email_date
    ws_gen["E3"].value = "Asunto del Correo:"
    ws_gen["F3"].value = latest_correo['subject']
    ws_gen["I3"].value = "Estado:"
    ws_gen["J3"].value = "ACTUALIZADO AL DÍA"

    # Encabezados de tabla (Fila 7)
    headers_gen = [
        "N°", "Sección Correo", "Categoría Hardware", "Marca",
        "N/P (Num. Parte)", "Producto / Descripción", "Costo COP",
        "IVA", "Fecha Ingreso", "Unidades Disponibles (HOY)", "Valor Total Stock (COP)"
    ]
    for col_idx, h_name in enumerate(headers_gen, start=1):
        ws_gen.cell(7, col_idx).value = h_name

    # Rellenar productos exactamente en el orden 1:1 del correo
    for idx, p in enumerate(today_products, start=1):
        target_row = 7 + idx
        ws_gen.cell(target_row, 1).value = idx
        ws_gen.cell(target_row, 2).value = p['seccion']
        ws_gen.cell(target_row, 3).value = p['categoria']
        ws_gen.cell(target_row, 4).value = p['marca']
        ws_gen.cell(target_row, 5).value = p['np']
        ws_gen.cell(target_row, 6).value = p['producto']
        ws_gen.cell(target_row, 7).value = p['costo']
        ws_gen.cell(target_row, 8).value = p['iva']
        ws_gen.cell(target_row, 9).value = p['fecha_ingreso']
        ws_gen.cell(target_row, 10).value = p['cantidad_actual']
        ws_gen.cell(target_row, 11).value = f"=G{target_row}*J{target_row}"

    # Limpiar filas sobrantes previas para que la hoja quede exactamente con el tamaño del correo
    last_product_row = 7 + len(today_products)
    for r in range(last_product_row + 1, last_product_row + 150):
        for c in range(1, 12):
            cell = ws_gen.cell(r, c)
            if type(cell).__name__ != 'MergedCell':
                cell.value = None

    log_message(f"✓ Hoja 'Generador Correo Inventario' sincronizada 1:1 con {len(today_products)} productos.")

    # 2. ACTUALIZAR HOJA 'MATRIZ INVENTARIO' EN SEGUNDO PLANO
    try:
        month_label = f"{latest_correo['month_str']} {latest_correo['year']}"
        row_map_month = {}
        for r in range(2, ws_matriz.max_row + 1):
            m_val = str(ws_matriz.cell(r, 1).value or "").strip()
            np_v = ws_matriz.cell(r, 4).value
            if m_val.upper() == month_label.upper() and np_v:
                row_map_month[str(np_v).strip()] = r

        for p in today_products:
            np_k = p['np'].strip()
            if np_k in row_map_month:
                t_row = row_map_month[np_k]
            else:
                t_row = ws_matriz.max_row + 1
                row_map_month[np_k] = t_row
                ws_matriz.cell(t_row, 1).value = month_label
                ws_matriz.cell(t_row, 2).value = p['categoria']
                ws_matriz.cell(t_row, 3).value = p['marca']
                ws_matriz.cell(t_row, 4).value = np_k
                ws_matriz.cell(t_row, 5).value = p['producto']
                ws_matriz.cell(t_row, 6).value = p['costo']
                ws_matriz.cell(t_row, 7).value = p['iva']
                ws_matriz.cell(t_row, 8).value = p['fecha_ingreso']
                ws_matriz.cell(t_row, 9).value = 3950.0

            ws_matriz.cell(t_row, 10).value = p['cantidad_actual']
            for d_key, d_qty in p['dias_stock'].items():
                d_num = int(d_key.split('_')[1])
                if 1 <= d_num <= 31:
                    ws_matriz.cell(t_row, 10 + d_num).value = d_qty
    except Exception as e:
        log_message(f"Aviso actualizando Matriz Inventario: {e}")

    # Guardar Excel
    try:
        wb.save(EXCEL_PATH)
        log_message(f"✓ Excel guardado exitosamente en OneDrive: {EXCEL_PATH}")
    except Exception as e:
        log_message(f"ERROR guardando Excel: {e}")
        return False

    # 3. Sincronizar base de datos JSON secundaria
    try:
        src_json = os.path.join(PROJECT_PATHS[0], "inventory_data.json")
        if os.path.exists(src_json):
            with open(src_json, 'r', encoding='utf-8') as f:
                existing_data = json.load(f)
            
            hist_prods = [p for p in existing_data.get('products', []) if p.get('mes') != month_label]
            new_month_prods = []
            base_id = len(hist_prods) + 1
            for idx, p in enumerate(today_products, start=base_id):
                item = dict(p)
                item['id'] = idx
                item['mes'] = month_label
                item['trm_ingreso'] = 3950.0
                new_month_prods.append(item)
            
            all_prods = hist_prods + new_month_prods
            payload = {
                "last_updated": datetime.now().isoformat(),
                "excel_path": EXCEL_PATH,
                "total_products": len(all_prods),
                "products": all_prods
            }
            for proj in PROJECT_PATHS:
                jp = os.path.join(proj, "inventory_data.json")
                try:
                    with open(jp, 'w', encoding='utf-8') as f:
                        json.dump(payload, f, ensure_ascii=False, indent=2)
                except Exception:
                    pass
    except Exception as e:
        pass

    total_unidades = sum(p['cantidad_actual'] for p in today_products)
    total_valor = sum(p['costo'] * p['cantidad_actual'] for p in today_products)
    log_message(f"Resumen: {len(today_products)} referencias | {total_unidades:,} unidades | ${total_valor:,.2f} COP")
    log_message("=== SINCRONIZACIÓN FINALIZADA CON ÉXITO ===")
    return True

def main():
    try:
        force_run = '--force' in sys.argv
        actualizar_excel_y_sistema(force=force_run)
    except Exception as e:
        err = traceback.format_exc()
        log_message(f"ERROR CRÍTICO INESPERADO:\n{err}")

if __name__ == '__main__':
    main()
